import os
import re
import secrets
import sqlite3
import time
import uuid
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, Request, Response, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from pydantic import BaseModel, Field, field_validator
from backend.config import PROJECT_ROOT
from .database import database
from .security import (COOKIE, SESSION_SECONDS, origin, secure_cookie, digest,
                       hash_password, check_password, check_origin, limit, current_user, require_user)
from .mail import send_link

router = APIRouter()

MAX_RESUME_FILE_BYTES = 5 * 1024 * 1024
MAX_RESUME_CHARACTERS = 12_000


def _resume_text(raw: bytes, suffix: str) -> str:
    """Extract résumé text without persisting the uploaded document."""
    if suffix == ".txt":
        text = raw.decode("utf-8-sig")
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw), strict=False)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            raise HTTPException(503, "PDF résumé support is not installed on this server.") from None
        except Exception:
            raise HTTPException(422, "This PDF has no readable text. Paste the relevant sections instead.") from None
    elif suffix == ".docx":
        try:
            from docx import Document
            document = Document(BytesIO(raw))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        except ImportError:
            raise HTTPException(503, "DOCX résumé support is not installed on this server.") from None
        except Exception:
            raise HTTPException(422, "This Word document has no readable text. Paste the relevant sections instead.") from None
    else:
        raise HTTPException(422, "Choose a .txt, .pdf or .docx résumé file.")
    if "\x00" in text:
        raise HTTPException(422, "This file has no readable résumé text. Paste the relevant sections instead.")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise HTTPException(422, "This file has no readable résumé text. Paste the relevant sections instead.")
    if len(text) > MAX_RESUME_CHARACTERS:
        raise HTTPException(422, "The résumé exceeds 12,000 characters. Import a shorter file or paste the relevant sections.")
    return text


class Credentials(BaseModel):
    email: str = Field(max_length=254)
    password: str = Field(min_length=12, max_length=128)

    @field_validator('email')
    @classmethod
    def email_address(cls, value):
        value = value.strip().lower()
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value):
            raise ValueError('Enter a valid email address.')
        return value


class Signup(Credentials):
    name: str = Field(min_length=1, max_length=80)


class EmailOnly(BaseModel):
    email: str = Field(max_length=254)


class LinkToken(BaseModel):
    token: str = Field(min_length=32, max_length=128)


class Reset(LinkToken):
    password: str = Field(min_length=12, max_length=128)


def auth_limit(request: Request):
    check_origin(request)
    limit('auth-ip:'+digest(request.client.host if request.client else 'unknown'), 20, 900)


def issue_link(user, purpose):
    token = secrets.token_urlsafe(32)
    with database() as db:
        db.execute('DELETE FROM links WHERE expires<?', (int(time.time()),))
        db.execute('INSERT INTO links VALUES (?,?,?,?)', (digest(token),user['id'],purpose,int(time.time())+1800))
    route = '/verify-email' if purpose == 'verify' else '/reset-password'
    link = origin()+route+'#token='+token
    try:
        dev_link = send_link(user['email'], purpose, link)
    except Exception:
        with database() as db:
            db.execute('DELETE FROM links WHERE token=?', (digest(token),))
        raise HTTPException(503, 'Email delivery is unavailable. Please try again later.') from None
    return dev_link


@router.get('/', include_in_schema=False)
def home(request: Request):
    return RedirectResponse('/interview' if current_user(request) else '/login', status_code=303)


@router.get('/login', response_class=HTMLResponse)
@router.get('/signup', response_class=HTMLResponse)
@router.get('/forgot-password', response_class=HTMLResponse)
@router.get('/reset-password', response_class=HTMLResponse)
@router.get('/verify-email', response_class=HTMLResponse)
def auth_page(request: Request):
    if request.url.path in ('/login','/signup') and current_user(request):
        return RedirectResponse('/interview', status_code=303)
    text = (PROJECT_ROOT/'templates/auth.html').read_text(encoding='utf-8')
    hint = 'Development email is saved in data/outbox (or your configured data directory).' if os.getenv('MAIL_MODE','file') == 'file' and os.getenv('APP_ENV','development') != 'production' else 'Check your inbox for the verification or reset link.'
    return HTMLResponse(text.replace('<!-- mail-hint -->',hint), headers={'Cache-Control':'no-store','Referrer-Policy':'no-referrer'})


@router.post('/auth/signup', dependencies=[Depends(auth_limit)])
def signup(data: Signup):
    limit('mail:'+digest(data.email), 3, 900)
    if not data.name.strip():
        raise HTTPException(422, 'Enter your name.')
    dev_link = None
    with database(readonly=True) as db:
        user = db.execute('SELECT * FROM users WHERE email=?', (data.email,)).fetchone()
    if not user:
        uid = str(uuid.uuid4())
        password = hash_password(data.password)
        try:
            with database() as db:
                db.execute('INSERT INTO users VALUES (?,?,?,?,0,?)', (uid,data.email,data.name.strip(),password,int(time.time())))
            user = {'id':uid,'email':data.email,'verified':0}
        except sqlite3.IntegrityError:
            user = None
    if user and not user['verified']:
        dev_link = issue_link(user, 'verify')
    response = {'message':'If this address is eligible, a verification link has been sent. Already registered? Sign in or reset your password.'}
    if dev_link:
        response['dev_link'] = dev_link
    return response


@router.post('/auth/login', dependencies=[Depends(auth_limit)])
def login(data: Credentials, response: Response, request: Request):
    limit('login-email:'+digest(data.email), 10, 900)
    with database(readonly=True) as db:
        user = db.execute('SELECT * FROM users WHERE email=?', (data.email,)).fetchone()
    if not user:
        hash_password(data.password)  # Comparable password-work for unknown accounts.
        raise HTTPException(401, 'Email or password is incorrect.')
    if not check_password(data.password, user['password']):
        raise HTTPException(401, 'Email or password is incorrect.')
    if not user['verified']:
        raise HTTPException(403, 'Verify your email before signing in. Use Resend verification below.')
    token = secrets.token_urlsafe(32)
    with database() as db:
        db.execute('BEGIN IMMEDIATE')
        if db.execute('SELECT 1 FROM account_suspensions WHERE user_id=?',(user['id'],)).fetchone():
            raise HTTPException(403, 'This account is suspended. Contact the site administrator.')
        db.execute('DELETE FROM logins WHERE expires<? OR token=?', (int(time.time()),digest(request.cookies.get(COOKIE,''))))
        db.execute('INSERT INTO logins VALUES (?,?,?)', (digest(token),user['id'],int(time.time())+SESSION_SECONDS))
    response.set_cookie(COOKIE, token, max_age=SESSION_SECONDS, httponly=True, secure=secure_cookie(), samesite='lax')
    response.headers['Cache-Control'] = 'no-store'
    return {'message':'Signed in.'}


@router.post('/auth/logout', dependencies=[Depends(check_origin)])
def logout(request: Request, response: Response):
    if current_user(request):
        require_user(request)
    with database() as db:
        db.execute('DELETE FROM logins WHERE token=?', (digest(request.cookies.get(COOKIE,'')),))
    response.delete_cookie(COOKIE, httponly=True, secure=secure_cookie(), samesite='lax')
    return {'message':'Signed out.'}


@router.get('/auth/me')
def me(user=Depends(require_user)):
    from backend.usage.access import is_admin
    return {**user, "is_admin": is_admin(user)}


@router.post('/api/account/resume-text', dependencies=[Depends(check_origin)])
def import_resume(file: UploadFile, user=Depends(require_user)):
    """Extract résumé text for the signed-in user; uploaded bytes are never saved."""
    filename = (file.filename or "").strip().lower()
    suffix = os.path.splitext(filename)[1]
    if suffix not in {".txt", ".pdf", ".docx"}:
        raise HTTPException(422, "Choose a .txt, .pdf or .docx résumé file.")
    raw = file.file.read(MAX_RESUME_FILE_BYTES + 1)
    if len(raw) > MAX_RESUME_FILE_BYTES:
        raise HTTPException(413, "The résumé file is larger than 5 MB.")
    try:
        text = _resume_text(raw, suffix)
    except UnicodeDecodeError:
        raise HTTPException(422, "This text file is not UTF-8 encoded. Save it as UTF-8 or use PDF/DOCX.") from None
    return {"text": text, "characters": len(text), "format": suffix[1:]}


@router.post('/auth/resend-verification', dependencies=[Depends(auth_limit)])
@router.post('/auth/forgot-password', dependencies=[Depends(auth_limit)])
def send_email(data: EmailOnly, request: Request):
    email = data.email.strip().lower()
    limit('mail:'+digest(email), 3, 900)
    dev_link = None
    with database(readonly=True) as db:
        user = db.execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
    purpose = 'verify' if request.url.path.endswith('resend-verification') else 'reset'
    if user and (purpose == 'reset' or not user['verified']):
        dev_link = issue_link(user,purpose)
    response = {'message':'If the account is eligible, an email has been sent.'}
    if dev_link:
        response['dev_link'] = dev_link
    return response


@router.post('/auth/verify-email', dependencies=[Depends(auth_limit)])
def verify(data: LinkToken):
    with database() as db:
        db.execute('BEGIN IMMEDIATE')
        link = db.execute('SELECT * FROM links WHERE token=? AND purpose=? AND expires>?', (digest(data.token),'verify',int(time.time()))).fetchone()
        if not link:
            raise HTTPException(400,'This link is invalid or expired. Request a new verification email.')
        db.execute('UPDATE users SET verified=1 WHERE id=?', (link['user_id'],))
        db.execute('DELETE FROM links WHERE user_id=? AND purpose=?',(link['user_id'],'verify'))
    return {'message':'Email verified. You can now sign in.'}


@router.post('/auth/reset-password', dependencies=[Depends(auth_limit)])
def reset(data: Reset):
    password = hash_password(data.password)
    with database() as db:
        db.execute('BEGIN IMMEDIATE')
        link = db.execute('SELECT * FROM links WHERE token=? AND purpose=? AND expires>?', (digest(data.token),'reset',int(time.time()))).fetchone()
        if not link:
            raise HTTPException(400,'This link is invalid or expired. Request a new password reset.')
        db.execute('UPDATE users SET password=? WHERE id=?',(password,link['user_id']))
        db.execute('DELETE FROM logins WHERE user_id=?',(link['user_id'],))
        db.execute('DELETE FROM links WHERE user_id=?',(link['user_id'],))
    return {'message':'Password updated. Sign in with your new password.'}
